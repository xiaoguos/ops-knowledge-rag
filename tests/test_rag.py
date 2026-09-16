import json
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from app.documents import split_markdown, parse_file
from app.retrieval import Embedder, bm25, rrf, query_route
from app.service import RagService


@pytest.fixture
def service(tmp_path):
    service = RagService(tmp_path)
    service.seed()
    return service


def test_structural_code_and_table_are_atomic():
    code = "```python\n" + "print('hello')\n" * 50 + "```"
    table = "| key | value |\n| --- | --- |\n| timeout | 3000 |"
    sections = split_markdown("# 配置\n\n前言\n" + code + "\n\n" + table, limit=100)
    assert any(code in section.text for section in sections)
    assert any(table in section.text for section in sections)
    assert all(section.heading == "配置" for section in sections)


def test_docx_preserves_paragraph_table_order():
    from docx import Document
    doc = Document()
    doc.add_heading("配置", 1)
    doc.add_paragraph("之前")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "timeout"
    table.cell(0, 1).text = "3000"
    doc.add_paragraph("之后")
    stream = BytesIO(); doc.save(stream)
    text = "\n".join(s.text for s in parse_file("manual.docx", stream.getvalue()))
    assert text.index("之前") < text.index("timeout") < text.index("之后")


def test_empty_pdf_explicitly_rejected():
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    stream = BytesIO(); writer.write(stream)
    with pytest.raises(ValueError, match="OCR"):
        parse_file("scan.pdf", stream.getvalue())


def test_version_and_department_filters_precede_retrieval(service):
    result = service.ask("订单查询接口每秒限流多少次", department="研发", version="2.0")
    assert not result["refused"]
    assert all(h["version"] == "2.0" and h["department"] == "研发" for h in result["hits"])
    assert "100" in result["answer"]
    assert service.ask("退款率怎么计算", department="研发")["refused"]


def test_unknown_question_refuses(service):
    assert service.ask("公司董事长的电话号码是多少")["refused"]


def test_rrf_merges_ranks_not_scores():
    ranking, scores = rrf([[0, 1], [1, 2]])
    assert ranking[0] == 1
    assert scores[1] == pytest.approx(1/62 + 1/61)


def test_identifier_route():
    assert query_route("PAY_TIMEOUT 怎么办")[0] == "bm25"
    assert query_route("服务变慢如何检查")[0] == "hybrid"


def test_bm25_empty_documents():
    assert bm25("query", []) == []


def test_atomic_update_removes_old_chunks(service):
    service.store.upsert("test", "x.md", "# 配置\n阈值 12".encode())
    old_ids = {r["id"] for r in service.store.chunks() if r["document_id"] == "test"}
    service.store.upsert("test", "x.md", "# 配置\n阈值 99".encode(), version="2")
    rows = [r for r in service.store.chunks() if r["document_id"] == "test"]
    assert not old_ids.intersection(r["id"] for r in rows)
    assert all("12" not in r["text"] for r in rows)
    assert not service.store.upsert("test", "x.md", "# 配置\n阈值 99".encode(), version="2")["changed"]


def test_failed_embedding_preserves_old_document(service, monkeypatch):
    original = service.store.documents()
    def fail(*args, **kwargs):
        raise ValueError("embedding unavailable")
    monkeypatch.setattr(service.embedder, "embed", fail)
    with pytest.raises(ValueError):
        service.store.upsert("payment-timeout", "new.md", b"new content")
    assert service.store.documents() == original


def test_delete_updates_all_retrieval_paths(service):
    assert service.store.delete("payment-timeout")
    for strategy in ["bm25", "vector", "hybrid"]:
        assert all(h["document_id"] != "payment-timeout" for h in service.search("PAY_TIMEOUT", strategy)["hits"])


def test_model_fingerprint_mismatch(service):
    service.embedder.fingerprint = "changed"
    with pytest.raises(ValueError, match="索引模型"):
        service.search("Redis")


def test_generation_invalid_citation_rejected(service, monkeypatch):
    class Response:
        def raise_for_status(self): pass
        def json(self): return {"choices": [{"message": {"content": json.dumps({"claims": [{"text": "wrong", "citation_id": "made-up", "quote": "fake"}]})}}]}
    monkeypatch.setattr("app.service.httpx.post", lambda *a, **k: Response())
    assert service.generate("question", [{"id": "real", "text": "actual evidence"}]) == []


def test_generation_matching_quote_accepts(service, monkeypatch):
    class Response:
        def raise_for_status(self): pass
        def json(self): return {"choices": [{"message": {"content": json.dumps({"claims": [{"text": "配置为 3000", "citation_id": "real", "quote": "3000"}]})}}]}
    monkeypatch.setattr("app.service.httpx.post", lambda *a, **k: Response())
    assert len(service.generate("question", [{"id": "real", "text": "timeout 3000"}])) == 1


def test_optional_rerank_without_key_degrades(service, monkeypatch):
    monkeypatch.delenv("COHERE_API_KEY", raising=False)
    result = service.search("Redis 内存", rerank=True)
    assert result["hits"] and "未配置" in result["warning"]
